import io
import re
import unicodedata
from typing import Dict, List, Tuple

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


st.set_page_config(
    page_title="Consolidador de Fretes",
    page_icon="📊",
    layout="wide",
)


OUTPUT_COLUMNS = [
    "EMISSÃO",
    "TRANSPORTADOR",
    "PLACA",
    "TIPOLOGIA",
    "PAGA",
    "ORIGEM",
    "DESTINO",
    "VENCIMENTO",
    "KM TOTAL",
    "ADC (KM",
    "FRETE SAÍDA",
    "PEDÁGIO",
    "FRANQUIA",
    "OUTROS CRÉDITOS",
    "DESCONTO",
    "ABASTECIMENTO",
    "FRETE SEM DESCONTO",
    "TOTAL COM DESCONTO",
    "O.S VIAG",
    "O.S FECH",
    "MOTIVO",
]


COLUMN_ALIASES = {
    "EMISSÃO": [
        "EMISSÃO",
        "EMISSAO",
    ],
    "TRANSPORTADOR": [
        "TRANSPORTADOR",
        "TRANSPORTADOR/MOTORISTA",
        "TRANSPORTADOR / MOTORISTA",
    ],
    "PLACA": [
        "PLACA",
    ],
    "TIPOLOGIA": [
        "TIPOLOGIA",
    ],
    "PAGA": [
        "PAGA",
    ],
    "ORIGEM": [
        "ORIGEM",
    ],
    "DESTINO": [
        "DESTINO",
    ],
    "VENCIMENTO": [
        "VENCIMENTO",
    ],
    "KM TOTAL": [
        "KM TOTAL",
    ],
    "ADC (KM": [
        "ADC (KM",
        "ADC KM",
        "ADC(KM",
    ],
    "FRETE SAÍDA": [
        "FRETE SAÍDA",
        "FRETE SAIDA",
    ],
    "PEDÁGIO": [
        "PEDÁGIO",
        "PEDAGIO",
    ],
    "FRANQUIA": [
        "FRANQUIA",
        "FRANQUINA",
    ],
    "OUTROS CRÉDITOS": [
        "OUTROS CRÉDITOS",
        "OUTROS CREDITOS",
    ],
    "DESCONTO": [
        "DESCONTO",
    ],
    "ABASTECIMENTO": [
        "ABASTECIMENTO",
        "COMBUSTÍVEL",
        "COMBUSTIVEL",
    ],
    "FRETE SEM DESCONTO": [
        "FRETE SEM DESCONTO",
    ],
    "TOTAL COM DESCONTO": [
        "TOTAL COM DESCONTO",
        "TOTAL COM DIFERENÇA / DESCONTO",
        "TOTAL COM DIFERENCA / DESCONTO",
        "TOTAL COM DIFERENÇA/DESCONTO",
        "TOTAL COM DIFERENCA/DESCONTO",
    ],
    "O.S VIAG": [
        "O.S VIAG",
        "O.S. VIAG",
        "OS VIAG",
    ],
    "O.S FECH": [
        "O.S FECH",
        "O.S. FECH",
        "OS FECH",
    ],
    "MOTIVO": [
        "MOTIVO",
        "OBSERVAÇÃO",
        "OBSERVACAO",
    ],
}


def normalize_text(value) -> str:
    """
    Normaliza somente para comparação de cabeçalhos.
    Não altera os valores existentes nas células.
    """
    if value is None:
        return ""

    text = str(value).replace("\xa0", " ")
    text = text.strip().lower()

    text = unicodedata.normalize("NFKD", text)
    text = "".join(
        char for char in text
        if not unicodedata.combining(char)
    )

    return re.sub(r"[^a-z0-9]+", "", text)


NORMALIZED_ALIASES = {
    target: [normalize_text(alias) for alias in aliases]
    for target, aliases in COLUMN_ALIASES.items()
}


def is_summary_sheet(sheet_name: str) -> bool:
    """
    Ignora abas cujo nome contenha a palavra 'resumo',
    independentemente de acento, espaços ou maiúsculas.
    """
    return "resumo" in normalize_text(sheet_name)


def has_any_value(row: pd.Series) -> bool:
    """
    Identifica se a linha original possui ao menos
    uma célula preenchida.
    """
    for value in row.tolist():
        if value is None:
            continue

        if isinstance(value, str) and value.strip() == "":
            continue

        try:
            if pd.isna(value):
                continue
        except (TypeError, ValueError):
            pass

        return True

    return False


def build_header_lookup(columns) -> Dict[str, object]:
    """
    Cria um mapa:
    cabeçalho normalizado -> cabeçalho original.

    Se houver cabeçalhos repetidos, mantém a primeira ocorrência.
    """
    lookup = {}

    for column in columns:
        normalized = normalize_text(column)

        if normalized and normalized not in lookup:
            lookup[normalized] = column

    return lookup


def find_source_column(
    header_lookup: Dict[str, object],
    target_column: str,
):
    """
    Procura qual coluna da aba corresponde à coluna
    esperada no arquivo final.
    """
    for alias in NORMALIZED_ALIASES[target_column]:
        if alias in header_lookup:
            return header_lookup[alias]

    return None


def detect_header_row(
    excel_file: pd.ExcelFile,
    sheet_name: str,
    max_rows: int = 10,
):
    """
    Procura o cabeçalho nas primeiras `max_rows` linhas da aba.

    A linha escolhida é aquela que contém a maior quantidade
    de cabeçalhos reconhecidos pelas regras de equivalência.

    Retorna:
    - índice da linha para usar no parâmetro header do pandas;
    - quantidade de campos reconhecidos.

    Se nenhuma linha tiver pelo menos 3 campos reconhecidos,
    retorna None.
    """

    preview = excel_file.parse(
        sheet_name=sheet_name,
        header=None,
        nrows=max_rows,
        dtype=object,
        keep_default_na=False,
    )

    best_row = None
    best_score = 0

    for row_index, row in preview.iterrows():

        normalized_values = {
            normalize_text(value)
            for value in row.tolist()
            if normalize_text(value)
        }

        matched_targets = set()

        for target_column, aliases in NORMALIZED_ALIASES.items():
            if any(
                alias in normalized_values
                for alias in aliases
            ):
                matched_targets.add(
                    target_column
                )

        score = len(matched_targets)

        if score > best_score:
            best_score = score
            best_row = int(row_index)

    # Segurança:
    # evita considerar como cabeçalho uma linha qualquer
    # que tenha apenas uma ou duas palavras coincidentes.
    if best_score < 3:
        return None, best_score

    return best_row, best_score



def is_blank_cell(value) -> bool:
    """
    Diz se uma célula está realmente vazia.

    Zero é considerado valor existente.
    """
    if value is None:
        return True

    if isinstance(value, str):
        return value.strip() == ""

    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def to_number(value) -> float:
    """
    Converte um valor para número somente para os cálculos
    de FRETE SEM DESCONTO e TOTAL COM DESCONTO.

    Aceita:
    - números do Excel;
    - valores como 1234,56;
    - valores como 1.234,56;
    - valores com R$;
    - valores negativos.

    Se a célula estiver vazia, considera 0 somente na conta.
    A célula original não é alterada.
    """
    if is_blank_cell(value):
        return 0.0

    if isinstance(value, bool):
        return float(value)

    if isinstance(value, (int, float)):
        return float(value)

    text = str(value).strip()
    text = text.replace("\xa0", " ")
    text = text.replace("R$", "")
    text = text.replace(" ", "")

    negative_parentheses = (
        text.startswith("(")
        and text.endswith(")")
    )

    if negative_parentheses:
        text = text[1:-1]

    # Mantém somente caracteres que podem compor um número.
    text = re.sub(r"[^0-9,.\-+]", "", text)

    if not text:
        return 0.0

    # Padrão brasileiro: 1.234,56
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "")
            text = text.replace(",", ".")
        else:
            # Ex.: 1,234.56
            text = text.replace(",", "")

    elif "," in text:
        # Ex.: 1234,56
        text = text.replace(".", "")
        text = text.replace(",", ".")

    elif "." in text:
        # Se houver vários pontos, considera-os separadores
        # de milhar, preservando o último apenas quando fizer
        # sentido como decimal.
        parts = text.split(".")

        if len(parts) > 2:
            if len(parts[-1]) in (1, 2):
                text = "".join(parts[:-1]) + "." + parts[-1]
            else:
                text = "".join(parts)

        elif len(parts) == 2:
            # No padrão brasileiro, "1.234" normalmente significa
            # mil duzentos e trinta e quatro.
            if len(parts[-1]) == 3 and parts[0].lstrip("+-").isdigit():
                text = "".join(parts)

    try:
        number = float(text)
    except ValueError:
        return 0.0

    if negative_parentheses:
        number *= -1

    return number


def clean_calculated_number(value: float):
    """
    Evita resíduos de ponto flutuante como 1199.9999999998.
    Mantém o resultado como número, não como fórmula.
    """
    rounded = round(float(value), 10)

    if rounded == int(rounded):
        return int(rounded)

    return rounded


def calculate_export_values(
    final_df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Recalcula APENAS:
    - FRETE SEM DESCONTO
    - TOTAL COM DESCONTO

    Os cálculos são feitos depois que todas as abas já foram
    consolidadas.

    Sem KM TOTAL:
        FRETE SEM DESCONTO =
            FRETE SAÍDA
            + PEDÁGIO
            + OUTROS CRÉDITOS

    Com KM TOTAL:
        FRETE SEM DESCONTO =
            FRETE SAÍDA
            + (ADC (KM * FRANQUIA)
            + PEDÁGIO
            + OUTROS CRÉDITOS

    Para todas as linhas:
        TOTAL COM DESCONTO =
            FRETE SEM DESCONTO
            + DESCONTO
            + ABASTECIMENTO

    Importante:
    - valores vazios entram como 0 somente na conta;
    - DESCONTO e ABASTECIMENTO são somados exatamente com
      o sinal que possuírem na planilha;
    - nenhuma fórmula é gravada no arquivo final.
    """

    if final_df.empty:
        return final_df

    final_df = final_df.copy()

    for row_index, row in final_df.iterrows():

        frete_saida = to_number(
            row.get("FRETE SAÍDA", "")
        )

        pedagio = to_number(
            row.get("PEDÁGIO", "")
        )

        outros_creditos = to_number(
            row.get("OUTROS CRÉDITOS", "")
        )

        desconto = to_number(
            row.get("DESCONTO", "")
        )

        abastecimento = to_number(
            row.get("ABASTECIMENTO", "")
        )

        km_total = row.get(
            "KM TOTAL",
            "",
        )

        tem_km_total = not is_blank_cell(
            km_total
        )

        if tem_km_total:

            adc_km = to_number(
                row.get("ADC (KM", "")
            )

            franquia = to_number(
                row.get("FRANQUIA", "")
            )

            frete_sem_desconto = (
                frete_saida
                + (adc_km * franquia)
                + pedagio
                + outros_creditos
            )

        else:

            frete_sem_desconto = (
                frete_saida
                + pedagio
                + outros_creditos
            )

        total_com_desconto = (
            frete_sem_desconto
            + desconto
            + abastecimento
        )

        final_df.at[
            row_index,
            "FRETE SEM DESCONTO",
        ] = clean_calculated_number(
            frete_sem_desconto
        )

        final_df.at[
            row_index,
            "TOTAL COM DESCONTO",
        ] = clean_calculated_number(
            total_com_desconto
        )

    return final_df


def process_workbook(
    uploaded_file,
) -> Tuple[pd.DataFrame, List[dict], List[str], List[str]]:
    """
    Percorre todas as abas e consolida os dados.

    Retorna:
    - DataFrame final;
    - relatório por aba;
    - abas de resumo ignoradas;
    - avisos encontrados.
    """

    # -------------------------------------------------------------
    # LER O EXCEL SOMENTE COM OS RESULTADOS DAS FÓRMULAS
    # -------------------------------------------------------------
    #
    # data_only=True faz com que:
    #   célula normal  -> copie o valor da célula
    #   célula fórmula -> copie somente o resultado salvo da fórmula
    #
    # Exemplo:
    #   fórmula original: =K10+L10
    #   valor exibido:    1250
    #   valor importado:  1250
    #
    # A fórmula NÃO é levada para a planilha final.
    # -------------------------------------------------------------

    original_bytes = uploaded_file.getvalue()

    values_workbook = load_workbook(
        io.BytesIO(original_bytes),
        data_only=True,
    )

    # Identifica abas ocultas antes do processamento.
    # hidden e veryHidden não serão processadas.
    hidden_sheet_names = {
        worksheet.title
        for worksheet in values_workbook.worksheets
        if worksheet.sheet_state != "visible"
    }

    # Salva uma cópia interna contendo apenas os valores já
    # calculados/salvos no arquivo original.
    values_buffer = io.BytesIO()

    values_workbook.save(
        values_buffer
    )

    values_buffer.seek(0)

    excel_file = pd.ExcelFile(
        values_buffer,
        engine="openpyxl",
    )

    consolidated_frames = []
    report = []
    ignored_sheets = []
    warnings = []

    for sheet_name in excel_file.sheet_names:

        # ---------------------------------------------------------
        # 1. IGNORAR ABAS OCULTAS E ABAS DE RESUMO
        # ---------------------------------------------------------

        if sheet_name in hidden_sheet_names:

            ignored_sheets.append(
                f"{sheet_name} (oculta)"
            )

            report.append(
                {
                    "Aba": sheet_name,
                    "Status": "Ignorada: aba oculta",
                    "Linhas importadas": 0,
                    "Colunas encontradas": 0,
                    "Colunas ausentes": 0,
                    "Ausentes": "-",
                }
            )

            continue

        if is_summary_sheet(sheet_name):

            ignored_sheets.append(
                f"{sheet_name} (resumo)"
            )

            continue

        # ---------------------------------------------------------
        # 2. LOCALIZAR E LER O CABEÇALHO
        # Procura automaticamente nas primeiras 10 linhas da aba.
        # ---------------------------------------------------------

        try:
            header_row, header_score = detect_header_row(
                excel_file=excel_file,
                sheet_name=sheet_name,
                max_rows=10,
            )

            if header_row is None:
                report.append(
                    {
                        "Aba": sheet_name,
                        "Status": (
                            "Ignorada: cabeçalho não encontrado "
                            "nas 10 primeiras linhas"
                        ),
                        "Linhas importadas": 0,
                        "Colunas encontradas": header_score,
                        "Colunas ausentes": len(OUTPUT_COLUMNS),
                        "Ausentes": "-",
                    }
                )

                warnings.append(
                    f'A aba "{sheet_name}" foi ignorada porque '
                    "não foi possível identificar com segurança "
                    "o cabeçalho nas 10 primeiras linhas."
                )

                continue

            df = excel_file.parse(
                sheet_name=sheet_name,
                header=header_row,
                dtype=object,
                keep_default_na=False,
            )

        except Exception as exc:
            report.append(
                {
                    "Aba": sheet_name,
                    "Status": f"Erro ao ler: {exc}",
                    "Linhas importadas": 0,
                    "Colunas encontradas": 0,
                    "Colunas ausentes": len(OUTPUT_COLUMNS),
                    "Ausentes": "-",
                }
            )
            continue

        # ---------------------------------------------------------
        # 3. REMOVER COLUNAS SEM CABEÇALHO
        # ---------------------------------------------------------

        valid_columns = [
            column
            for column in df.columns
            if normalize_text(column)
            and not normalize_text(column).startswith("unnamed")
        ]

        df = df.loc[:, valid_columns]

        # ---------------------------------------------------------
        # 4. REMOVER SOMENTE LINHAS TOTALMENTE VAZIAS
        # ---------------------------------------------------------

        if not df.empty:
            non_empty_mask = df.apply(
                has_any_value,
                axis=1,
            )

            df = df.loc[non_empty_mask].copy()

        # ---------------------------------------------------------
        # 5. CRIAR MAPA DOS CABEÇALHOS ENCONTRADOS
        # ---------------------------------------------------------

        header_lookup = build_header_lookup(
            df.columns
        )

        column_mapping = {}
        missing_columns = []

        for target_column in OUTPUT_COLUMNS:

            source_column = find_source_column(
                header_lookup,
                target_column,
            )

            if source_column is None:
                missing_columns.append(
                    target_column
                )
            else:
                column_mapping[
                    target_column
                ] = source_column

        # ---------------------------------------------------------
        # 6. SE NÃO HOUVER NENHUMA COLUNA CONHECIDA,
        # PROVAVELMENTE NÃO É UMA ABA DE DADOS
        # ---------------------------------------------------------

        if not column_mapping:

            report.append(
                {
                    "Aba": sheet_name,
                    "Status": (
                        "Ignorada: nenhum "
                        "cabeçalho reconhecido"
                    ),
                    "Linhas importadas": 0,
                    "Colunas encontradas": 0,
                    "Colunas ausentes": len(
                        OUTPUT_COLUMNS
                    ),
                    "Ausentes": ", ".join(
                        OUTPUT_COLUMNS
                    ),
                }
            )

            warnings.append(
                f'A aba "{sheet_name}" não possui '
                "nenhum dos cabeçalhos esperados "
                "e não foi importada."
            )

            continue

        # ---------------------------------------------------------
        # 7. CRIAR A ESTRUTURA PADRÃO DA PLANILHA FINAL
        # ---------------------------------------------------------

        output_df = pd.DataFrame(
            index=df.index
        )

        for target_column in OUTPUT_COLUMNS:

            source_column = column_mapping.get(
                target_column
            )

            if source_column is None:

                # Se a coluna não existir nessa aba,
                # deixa em branco.
                output_df[target_column] = ""

            else:

                # Copia exatamente o valor encontrado.
                # Não faz conta, não preenche zero,
                # não altera nome, placa ou valor.
                output_df[target_column] = df[
                    source_column
                ]

        output_df.reset_index(
            drop=True,
            inplace=True,
        )

        consolidated_frames.append(
            output_df
        )

        # ---------------------------------------------------------
        # 8. DETECTAR POSSÍVEL "TIPOLOGIA PAGA"
        # ---------------------------------------------------------

        normalized_headers = {
            normalize_text(column)
            for column in df.columns
        }

        if (
            normalize_text("TIPOLOGIA PAGA")
            in normalized_headers
            and normalize_text("TIPOLOGIA")
            not in normalized_headers
            and normalize_text("PAGA")
            not in normalized_headers
        ):

            warnings.append(
                f'A aba "{sheet_name}" possui '
                'uma coluna literal "TIPOLOGIA PAGA". '
                "Ela não foi dividida automaticamente "
                "entre TIPOLOGIA e PAGA para evitar "
                "alteração ou interpretação indevida "
                "dos dados."
            )

        # ---------------------------------------------------------
        # 9. RELATÓRIO DA ABA
        # ---------------------------------------------------------

        report.append(
            {
                "Aba": sheet_name,
                "Status": "Processada",
                "Linhas importadas": len(
                    output_df
                ),
                "Colunas encontradas": len(
                    column_mapping
                ),
                "Colunas ausentes": len(
                    missing_columns
                ),
                "Ausentes": (
                    ", ".join(missing_columns)
                    if missing_columns
                    else "-"
                ),
            }
        )

    # -------------------------------------------------------------
    # 10. JUNTAR TODAS AS ABAS
    # -------------------------------------------------------------

    if consolidated_frames:

        final_df = pd.concat(
            consolidated_frames,
            ignore_index=True,
        )

        final_df = final_df[
            OUTPUT_COLUMNS
        ]

        # ---------------------------------------------------------
        # 11. RECALCULAR OS VALORES DEPOIS DA CONSOLIDAÇÃO
        # ---------------------------------------------------------
        #
        # Somente estas duas colunas são recalculadas:
        # - FRETE SEM DESCONTO
        # - TOTAL COM DESCONTO
        #
        # O resultado é gravado como número, nunca como fórmula.
        # ---------------------------------------------------------

        final_df = calculate_export_values(
            final_df
        )

    else:

        final_df = pd.DataFrame(
            columns=OUTPUT_COLUMNS
        )

    return (
        final_df,
        report,
        ignored_sheets,
        warnings,
    )


def motivo_indica_abastecimento(value) -> bool:
    """
    Retorna True quando o campo MOTIVO contém
    'abastecimento' ou 'combustível', ignorando
    maiúsculas/minúsculas, acentos e pontuação.
    """
    normalized = normalize_text(value)

    return (
        "abastecimento" in normalized
        or "combustivel" in normalized
    )


def dataframe_to_excel_bytes(
    df: pd.DataFrame,
) -> bytes:
    """
    Cria o arquivo Excel final em memória.
    """

    output = io.BytesIO()

    # -------------------------------------------------------------
    # 1. ESCREVER OS DADOS
    # -------------------------------------------------------------

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:

        df.to_excel(
            writer,
            index=False,
            sheet_name="Resumo Geral",
        )

    output.seek(0)

    # -------------------------------------------------------------
    # 2. FORMATAR SOMENTE A APRESENTAÇÃO
    # -------------------------------------------------------------

    workbook = load_workbook(
        output
    )

    worksheet = workbook[
        "Resumo Geral"
    ]

    # Congelar cabeçalho
    worksheet.freeze_panes = "A2"

    # Filtro automático
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    # Formatação do cabeçalho
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for cell in worksheet[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # -------------------------------------------------------------
    # DESTACAR MOTIVO QUANDO HOUVER ABASTECIMENTO SEM JUSTIFICATIVA
    # -------------------------------------------------------------
    #
    # Regra:
    # - Se ABASTECIMENTO tiver um número;
    # - e MOTIVO estiver vazio OU não contiver
    #   "abastecimento" ou "combustível";
    # - pintar somente a célula MOTIVO de amarelo.
    #
    # Nenhum valor é alterado.
    # -------------------------------------------------------------

    abastecimento_col_idx = OUTPUT_COLUMNS.index(
        "ABASTECIMENTO"
    ) + 1

    motivo_col_idx = OUTPUT_COLUMNS.index(
        "MOTIVO"
    ) + 1

    yellow_fill = PatternFill(
        fill_type="solid",
        fgColor="FFFF00",
    )

    for row_idx in range(
        2,
        worksheet.max_row + 1,
    ):

        abastecimento_value = worksheet.cell(
            row=row_idx,
            column=abastecimento_col_idx,
        ).value

        motivo_value = worksheet.cell(
            row=row_idx,
            column=motivo_col_idx,
        ).value

        abastecimento_tem_numero = False

        if abastecimento_value is not None:

            if isinstance(
                abastecimento_value,
                (int, float),
            ):
                abastecimento_tem_numero = True

            elif isinstance(
                abastecimento_value,
                str,
            ):
                texto_abastecimento = (
                    abastecimento_value
                    .replace("R$", "")
                    .replace(" ", "")
                    .strip()
                )

                if texto_abastecimento:
                    abastecimento_tem_numero = bool(
                        re.search(
                            r"\d",
                            texto_abastecimento,
                        )
                    )

        if (
            abastecimento_tem_numero
            and not motivo_indica_abastecimento(
                motivo_value
            )
        ):
            worksheet.cell(
                row=row_idx,
                column=motivo_col_idx,
            ).fill = yellow_fill

    # -------------------------------------------------------------
    # 3. AJUSTAR LARGURA DAS COLUNAS
    # Sem alterar os dados.
    # -------------------------------------------------------------

    for col_idx, column_name in enumerate(
        OUTPUT_COLUMNS,
        start=1,
    ):

        max_length = len(
            column_name
        )

        # Para desempenho, verifica no máximo
        # as primeiras 300 linhas.
        for row_idx in range(
            2,
            min(
                worksheet.max_row,
                300,
            )
            + 1,
        ):

            value = worksheet.cell(
                row=row_idx,
                column=col_idx,
            ).value

            if value is not None:

                max_length = max(
                    max_length,
                    len(str(value)),
                )

        width = min(
            max(
                max_length + 2,
                12,
            ),
            35,
        )

        worksheet.column_dimensions[
            get_column_letter(col_idx)
        ].width = width

    # -------------------------------------------------------------
    # 4. SALVAR O ARQUIVO FINAL
    # -------------------------------------------------------------

    final_output = io.BytesIO()

    workbook.save(
        final_output
    )

    final_output.seek(0)

    return final_output.getvalue()



# =================================================================
# ETAPA 2 - CONFERÊNCIA CONTÁBIL
# =================================================================

ACCOUNTING_ALIASES = {
    "DATA": [
        "DATA",
    ],
    "HISTORICO": [
        "HISTORICO",
        "HISTÓRICO",
    ],
    "CONTA": [
        "CONTA",
    ],
    "CODIGO": [
        "CODIGO",
        "CÓDIGO",
    ],
    "DESCRIÇÃO": [
        "DESCRIÇÃO",
        "DESCRICAO",
    ],
    "C.CUSTO": [
        "C.CUSTO",
        "C CUSTO",
        "CCUSTO",
        "CENTRO DE CUSTO",
    ],
    "DÉBITO": [
        "DÉBITO",
        "DEBITO",
    ],
    "CRÉDITO": [
        "CRÉDITO",
        "CREDITO",
    ],
    "SALDO": [
        "SALDO",
    ],
}

NORMALIZED_ACCOUNTING_ALIASES = {
    target: [
        normalize_text(alias)
        for alias in aliases
    ]
    for target, aliases in ACCOUNTING_ALIASES.items()
}


def accounting_target_for_header(value):
    """
    Identifica a qual coluna contábil padrão um cabeçalho corresponde.
    """
    normalized = normalize_text(value)

    if not normalized:
        return None

    for target, aliases in NORMALIZED_ACCOUNTING_ALIASES.items():
        if normalized in aliases:
            return target

    return None


def detect_accounting_header_row(
    excel_file: pd.ExcelFile,
    sheet_name: str,
    max_rows: int = 10,
):
    """
    Procura nas primeiras 10 linhas a linha que contém
    o cabeçalho da planilha contábil.

    Para ser aceita, a linha precisa conter HISTORICO e
    pelo menos dois entre DÉBITO, CRÉDITO e SALDO.
    """

    preview = excel_file.parse(
        sheet_name=sheet_name,
        header=None,
        nrows=max_rows,
        dtype=object,
        keep_default_na=False,
    )

    best_row = None
    best_score = 0

    for row_index, row in preview.iterrows():

        found = set()

        for value in row.tolist():
            target = accounting_target_for_header(
                value
            )

            if target:
                found.add(target)

        financial_matches = len(
            found.intersection(
                {
                    "DÉBITO",
                    "CRÉDITO",
                    "SALDO",
                }
            )
        )

        score = len(found)

        if (
            "HISTORICO" in found
            and financial_matches >= 2
            and score > best_score
        ):
            best_row = int(row_index)
            best_score = score

    return best_row, best_score


def build_accounting_column_map(columns):
    """
    Localiza as colunas necessárias na planilha contábil.

    Se houver mais de uma coluna com o mesmo significado,
    a aba é considerada ambígua e não é processada.
    """

    matches = {
        target: []
        for target in ACCOUNTING_ALIASES
    }

    for column in columns:

        target = accounting_target_for_header(
            column
        )

        if target:
            matches[target].append(
                column
            )

    mapped = {}
    ambiguous = {}

    for target, candidates in matches.items():

        if len(candidates) == 1:
            mapped[target] = candidates[0]

        elif len(candidates) > 1:
            ambiguous[target] = candidates

    return mapped, ambiguous


def normalize_os_id(value):
    """
    Normaliza O.S. para comparação sem alterar o valor exibido.

    Exemplos:
    13929     -> "13929"
    13929.0   -> "13929"
    "13929"   -> "13929"
    """

    if is_blank_cell(value):
        return ""

    if isinstance(value, bool):
        return ""

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))

        text = str(value).strip()
    else:
        text = str(value).strip()

    text = text.replace("\xa0", " ")

    match = re.fullmatch(
        r"(\d+)(?:[.,]0+)?",
        text,
    )

    if match:
        return match.group(1)

    # Evita interpretar textos arbitrários da coluna O.S.
    return ""


def extract_os_from_historico(value):
    """
    Só considera lançamento de O.S. quando HISTORICO começa com:

        OS.3 13929 ...

    Também tolera a grafia O.S.3, mantendo a mesma regra:
    precisa começar pelo marcador e depois trazer uma numeração.
    """

    if is_blank_cell(value):
        return ""

    text = str(value).strip()

    match = re.match(
        r"^O\.?S\.?3\s+(\d+)",
        text,
        flags=re.IGNORECASE,
    )

    if not match:
        return ""

    return match.group(1)


def money_equal(value_a, value_b) -> bool:
    """
    Compara valores monetários até 2 casas decimais.
    """
    return round(
        to_number(value_a),
        2,
    ) == round(
        to_number(value_b),
        2,
    )


def format_money_br(value) -> str:
    """
    Formata um número para texto brasileiro apenas nas
    mensagens de divergência.
    """
    number = round(
        to_number(value),
        2,
    )

    formatted = f"{number:,.2f}"

    return (
        formatted
        .replace(",", "X")
        .replace(".", ",")
        .replace("X", ".")
    )


def append_observation(current_value, new_text):
    """
    Preserva qualquer observação já existente e acrescenta
    a observação da conferência sem apagar conteúdo anterior.
    """

    current = (
        ""
        if is_blank_cell(current_value)
        else str(current_value).strip()
    )

    new_text = (
        ""
        if new_text is None
        else str(new_text).strip()
    )

    if not new_text:
        return current

    if not current:
        return new_text

    if new_text in current:
        return current

    return current + " | " + new_text


def prepare_consolidated_lookup(
    consolidated_df: pd.DataFrame,
):
    """
    Cria índice por O.S VIAG.

    Uma O.S. com mais de uma linha no consolidado é mantida
    como duplicada para impedir escolha automática.
    """

    lookup = {}

    if consolidated_df is None or consolidated_df.empty:
        return lookup

    required = [
        "O.S VIAG",
        "VENCIMENTO",
        "FRETE SEM DESCONTO",
        "DESCONTO",
        "ABASTECIMENTO",
        "TOTAL COM DESCONTO",
    ]

    missing = [
        column
        for column in required
        if column not in consolidated_df.columns
    ]

    if missing:
        raise ValueError(
            "O consolidado não possui as colunas necessárias: "
            + ", ".join(missing)
        )

    for row_index, row in consolidated_df.iterrows():

        os_id = normalize_os_id(
            row.get(
                "O.S VIAG",
                "",
            )
        )

        if not os_id:
            continue

        lookup.setdefault(
            os_id,
            [],
        ).append(
            {
                "row_index": row_index,
                "VENCIMENTO": row.get(
                    "VENCIMENTO",
                    "",
                ),
                "FRETE SEM DESCONTO": row.get(
                    "FRETE SEM DESCONTO",
                    "",
                ),
                "DESCONTO": row.get(
                    "DESCONTO",
                    "",
                ),
                "ABASTECIMENTO": row.get(
                    "ABASTECIMENTO",
                    "",
                ),
                "TOTAL COM DESCONTO": row.get(
                    "TOTAL COM DESCONTO",
                    "",
                ),
            }
        )

    return lookup


def load_consolidated_from_upload(
    uploaded_file,
) -> pd.DataFrame:
    """
    Carrega um Resumo Geral já gerado anteriormente.

    Procura o cabeçalho nas primeiras 10 linhas e busca uma
    aba visível que contenha as colunas essenciais.
    """

    original_bytes = uploaded_file.getvalue()

    wb_values = load_workbook(
        io.BytesIO(original_bytes),
        data_only=True,
    )

    visible_sheets = {
        ws.title
        for ws in wb_values.worksheets
        if ws.sheet_state == "visible"
    }

    buffer = io.BytesIO()
    wb_values.save(buffer)
    buffer.seek(0)

    excel_file = pd.ExcelFile(
        buffer,
        engine="openpyxl",
    )

    # Dá preferência à aba "Resumo Geral".
    ordered_sheets = sorted(
        excel_file.sheet_names,
        key=lambda name: (
            normalize_text(name)
            != normalize_text("Resumo Geral"),
            excel_file.sheet_names.index(name),
        ),
    )

    for sheet_name in ordered_sheets:

        if sheet_name not in visible_sheets:
            continue

        header_row, _ = detect_header_row(
            excel_file=excel_file,
            sheet_name=sheet_name,
            max_rows=10,
        )

        if header_row is None:
            continue

        df = excel_file.parse(
            sheet_name=sheet_name,
            header=header_row,
            dtype=object,
            keep_default_na=False,
        )

        header_lookup = build_header_lookup(
            df.columns
        )

        mapped = {}

        for target_column in OUTPUT_COLUMNS:

            source_column = find_source_column(
                header_lookup,
                target_column,
            )

            if source_column is not None:
                mapped[target_column] = (
                    source_column
                )

        essential = {
            "O.S VIAG",
            "VENCIMENTO",
            "FRETE SEM DESCONTO",
            "DESCONTO",
            "ABASTECIMENTO",
            "TOTAL COM DESCONTO",
        }

        if not essential.issubset(
            mapped.keys()
        ):
            continue

        result = pd.DataFrame(
            index=df.index
        )

        for target_column in OUTPUT_COLUMNS:

            source = mapped.get(
                target_column
            )

            if source is None:
                result[
                    target_column
                ] = ""
            else:
                result[
                    target_column
                ] = df[source]

        result.reset_index(
            drop=True,
            inplace=True,
        )

        return result[
            OUTPUT_COLUMNS
        ]

    raise ValueError(
        "Não encontrei uma aba visível do consolidado "
        "com O.S VIAG, VENCIMENTO, FRETE SEM DESCONTO, "
        "DESCONTO, ABASTECIMENTO e TOTAL COM DESCONTO."
    )


def values_only_excel_file(
    uploaded_file,
):
    """
    Retorna:
    - pd.ExcelFile lendo somente resultados salvos de fórmulas;
    - conjunto de abas ocultas.
    """

    original_bytes = uploaded_file.getvalue()

    wb_values = load_workbook(
        io.BytesIO(original_bytes),
        data_only=True,
    )

    hidden_sheets = {
        ws.title
        for ws in wb_values.worksheets
        if ws.sheet_state != "visible"
    }

    buffer = io.BytesIO()
    wb_values.save(buffer)
    buffer.seek(0)

    return (
        pd.ExcelFile(
            buffer,
            engine="openpyxl",
        ),
        hidden_sheets,
    )


def process_accounting_sheet(
    df: pd.DataFrame,
    column_map: dict,
    consolidated_lookup: dict,
):
    """
    Faz a conferência de uma aba contábil.

    Retorna:
    - DataFrame final;
    - estilos por índice de linha;
    - estatísticas.
    """

    original_columns = list(
        df.columns
    )

    # -------------------------------------------------------------
    # DESCOBRIR OU CRIAR COLUNAS NOVAS
    # -------------------------------------------------------------

    historico_column = column_map[
        "HISTORICO"
    ]

    debit_column = column_map.get(
        "DÉBITO"
    )

    credit_column = column_map.get(
        "CRÉDITO"
    )

    saldo_column = column_map.get(
        "SALDO"
    )

    # CONFERÊNCIA sempre será a primeira coluna.
    conference_existing = None

    for column in original_columns:
        if normalize_text(
            column
        ) == normalize_text(
            "CONFERÊNCIA"
        ):
            conference_existing = column
            break

    if conference_existing is None:
        df.insert(
            0,
            "CONFERÊNCIA",
            "",
        )
        conference_column = "CONFERÊNCIA"
    else:
        conference_column = (
            conference_existing
        )

        columns_order = [
            conference_column
        ] + [
            column
            for column in df.columns
            if column != conference_column
        ]

        df = df[
            columns_order
        ].copy()

    # VENCIMENTO preferencialmente logo após HISTORICO.
    vencimento_existing = None

    for column in df.columns:
        if normalize_text(
            column
        ) == normalize_text(
            "VENCIMENTO"
        ):
            vencimento_existing = column
            break

    if vencimento_existing is None:

        historico_position = (
            list(df.columns).index(
                historico_column
            )
        )

        df.insert(
            historico_position + 1,
            "VENCIMENTO",
            "",
        )

        vencimento_column = (
            "VENCIMENTO"
        )

    else:
        vencimento_column = (
            vencimento_existing
        )

    # OBSERVAÇÃO ao final, se ainda não existir.
    observation_existing = None

    for column in df.columns:

        if normalize_text(
            column
        ) in {
            normalize_text("OBSERVAÇÃO"),
            normalize_text("OBSERVACAO"),
        }:
            observation_existing = column
            break

    if observation_existing is None:

        df[
            "OBSERVAÇÃO"
        ] = ""

        observation_column = (
            "OBSERVAÇÃO"
        )

    else:
        observation_column = (
            observation_existing
        )

    # -------------------------------------------------------------
    # IDENTIFICAR O.S. POR LINHA
    # -------------------------------------------------------------

    os_by_index = {}

    for row_index, value in df[
        historico_column
    ].items():

        os_by_index[
            row_index
        ] = extract_os_from_historico(
            value
        )

    groups = {}

    for row_index, os_id in os_by_index.items():

        if os_id:
            groups.setdefault(
                os_id,
                [],
            ).append(
                row_index
            )

    # Estilos:
    # row_status: "red" ou None
    # observation_fill: "yellow" / "blue" / None
    row_styles = {
        index: {
            "row_fill": None,
            "observation_fill": None,
        }
        for index in df.index
    }

    stats = {
        "os_analisadas": 0,
        "conferidas": 0,
        "nao_conferidas": 0,
        "linhas_nao_os": sum(
            1
            for value in os_by_index.values()
            if not value
        ),
    }

    # -------------------------------------------------------------
    # ANALISAR CADA GRUPO DE O.S.
    # -------------------------------------------------------------

    for os_id, indexes in groups.items():

        stats[
            "os_analisadas"
        ] += 1

        matches = consolidated_lookup.get(
            os_id,
            [],
        )

        # ---------------------------------------------------------
        # O.S. NÃO ENCONTRADA
        # ---------------------------------------------------------

        if len(matches) == 0:

            note = (
                f"O.S. {os_id} não encontrada "
                "na planilha consolidada."
            )

            for index in indexes:

                df.at[
                    index,
                    conference_column,
                ] = "NÃO CONFERIDO"

                df.at[
                    index,
                    observation_column,
                ] = append_observation(
                    df.at[
                        index,
                        observation_column,
                    ],
                    note,
                )

                row_styles[
                    index
                ][
                    "row_fill"
                ] = "red"

            stats[
                "nao_conferidas"
            ] += 1

            continue

        # ---------------------------------------------------------
        # O.S. DUPLICADA NO CONSOLIDADO
        # ---------------------------------------------------------

        if len(matches) > 1:

            note = (
                f"O.S. {os_id} aparece {len(matches)} vezes "
                "na planilha consolidada. "
                "Conferência automática não realizada."
            )

            for index in indexes:

                df.at[
                    index,
                    conference_column,
                ] = "NÃO CONFERIDO"

                df.at[
                    index,
                    observation_column,
                ] = append_observation(
                    df.at[
                        index,
                        observation_column,
                    ],
                    note,
                )

                row_styles[
                    index
                ][
                    "row_fill"
                ] = "red"

            stats[
                "nao_conferidas"
            ] += 1

            continue

        consolidated = matches[0]

        # Vencimento para todas as linhas da O.S.
        for index in indexes:

            df.at[
                index,
                vencimento_column,
            ] = consolidated[
                "VENCIMENTO"
            ]

        # ---------------------------------------------------------
        # TOTAIS CONTÁBEIS DA O.S.
        # ---------------------------------------------------------

        debit_total = sum(
            to_number(
                df.at[
                    index,
                    debit_column,
                ]
            )
            for index in indexes
        )

        credit_total = sum(
            to_number(
                df.at[
                    index,
                    credit_column,
                ]
            )
            for index in indexes
        )

        saldo_total = sum(
            to_number(
                df.at[
                    index,
                    saldo_column,
                ]
            )
            for index in indexes
        )

        expected_debit = to_number(
            consolidated[
                "FRETE SEM DESCONTO"
            ]
        )

        abastecimento_value = abs(
            to_number(
                consolidated[
                    "ABASTECIMENTO"
                ]
            )
        )

        desconto_value = abs(
            to_number(
                consolidated[
                    "DESCONTO"
                ]
            )
        )

        expected_credit = (
            abastecimento_value
            + desconto_value
        )

        expected_saldo = to_number(
            consolidated[
                "TOTAL COM DESCONTO"
            ]
        )

        debit_ok = money_equal(
            debit_total,
            expected_debit,
        )

        credit_ok = money_equal(
            credit_total,
            expected_credit,
        )

        saldo_ok = money_equal(
            saldo_total,
            expected_saldo,
        )

        # ---------------------------------------------------------
        # CLASSIFICAR LINHAS DE CRÉDITO
        # ---------------------------------------------------------

        ambiguous_credit_classification = False

        for index in indexes:

            credit_value = to_number(
                df.at[
                    index,
                    credit_column,
                ]
            )

            if round(
                abs(credit_value),
                2,
            ) == 0:
                continue

            matches_abastecimento = (
                abastecimento_value > 0
                and money_equal(
                    abs(credit_value),
                    abastecimento_value,
                )
            )

            matches_desconto = (
                desconto_value > 0
                and money_equal(
                    abs(credit_value),
                    desconto_value,
                )
            )

            if (
                matches_abastecimento
                and matches_desconto
            ):

                ambiguous_credit_classification = True

                df.at[
                    index,
                    observation_column,
                ] = append_observation(
                    df.at[
                        index,
                        observation_column,
                    ],
                    (
                        "Crédito coincide com os valores de "
                        "ABASTECIMENTO e DESCONTO. "
                        "Verificar classificação manualmente."
                    ),
                )

            elif matches_abastecimento:

                df.at[
                    index,
                    observation_column,
                ] = append_observation(
                    df.at[
                        index,
                        observation_column,
                    ],
                    (
                        "Desconto de Combustível. "
                        "Reclassificar para conta de "
                        "adiantamento de combustível."
                    ),
                )

                row_styles[
                    index
                ][
                    "observation_fill"
                ] = "yellow"

            elif matches_desconto:

                df.at[
                    index,
                    observation_column,
                ] = append_observation(
                    df.at[
                        index,
                        observation_column,
                    ],
                    (
                        "Reclassificar para desconto de Avarias."
                    ),
                )

                row_styles[
                    index
                ][
                    "observation_fill"
                ] = "blue"

        # Se um crédito individual é ambíguo, não vamos
        # considerar a O.S. totalmente conferida.
        classification_ok = (
            not ambiguous_credit_classification
        )

        all_ok = (
            debit_ok
            and credit_ok
            and saldo_ok
            and classification_ok
        )

        if all_ok:

            for index in indexes:

                df.at[
                    index,
                    conference_column,
                ] = "CONFERIDO"

            stats[
                "conferidas"
            ] += 1

        else:

            divergence_parts = []

            if not debit_ok:

                divergence_parts.append(
                    "Débito divergente: "
                    f"Contabilidade {format_money_br(debit_total)} | "
                    f"Consolidado {format_money_br(expected_debit)} | "
                    f"Diferença {format_money_br(debit_total - expected_debit)}."
                )

            if not credit_ok:

                divergence_parts.append(
                    "Crédito divergente: "
                    f"Contabilidade {format_money_br(credit_total)} | "
                    f"Esperado {format_money_br(expected_credit)} | "
                    f"Diferença {format_money_br(credit_total - expected_credit)}."
                )

            if not saldo_ok:

                divergence_parts.append(
                    "Saldo divergente: "
                    f"Contabilidade {format_money_br(saldo_total)} | "
                    f"Consolidado {format_money_br(expected_saldo)} | "
                    f"Diferença {format_money_br(saldo_total - expected_saldo)}."
                )

            if not classification_ok:

                divergence_parts.append(
                    "Classificação do crédito ambígua."
                )

            divergence_note = " ".join(
                divergence_parts
            )

            for index in indexes:

                df.at[
                    index,
                    conference_column,
                ] = "NÃO CONFERIDO"

                df.at[
                    index,
                    observation_column,
                ] = append_observation(
                    df.at[
                        index,
                        observation_column,
                    ],
                    divergence_note,
                )

                row_styles[
                    index
                ][
                    "row_fill"
                ] = "red"

            stats[
                "nao_conferidas"
            ] += 1

    return (
        df,
        row_styles,
        stats,
        {
            "conference_column": conference_column,
            "vencimento_column": vencimento_column,
            "observation_column": observation_column,
            "debit_column": debit_column,
            "credit_column": credit_column,
            "saldo_column": saldo_column,
        },
    )


def style_accounting_output_sheet(
    ws,
    df: pd.DataFrame,
    row_styles: dict,
    metadata: dict,
):
    """
    Formata a aba contábil final.

    Linha 1: subtotais.
    Linha 2: cabeçalhos.
    Linha 3 em diante: dados.
    """

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    subtotal_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    subtotal_font = Font(
        bold=True,
    )

    red_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC",
    )

    yellow_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    blue_fill = PatternFill(
        fill_type="solid",
        fgColor="CFE2F3",
    )

    # -------------------------------------------------------------
    # CABEÇALHO
    # -------------------------------------------------------------

    for cell in ws[2]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # -------------------------------------------------------------
    # SUBTOTAIS NA LINHA 1
    # -------------------------------------------------------------

    ws.cell(
        row=1,
        column=1,
        value="SUBTOTAL GERAL",
    )

    ws.cell(
        row=1,
        column=1,
    ).fill = subtotal_fill

    ws.cell(
        row=1,
        column=1,
    ).font = subtotal_font

    for logical_name in [
        "debit_column",
        "credit_column",
        "saldo_column",
    ]:

        column_name = metadata.get(
            logical_name
        )

        if column_name is None:
            continue

        column_index = (
            list(df.columns).index(
                column_name
            )
            + 1
        )

        total = sum(
            to_number(value)
            for value in df[
                column_name
            ].tolist()
        )

        cell = ws.cell(
            row=1,
            column=column_index,
            value=clean_calculated_number(
                total
            ),
        )

        cell.fill = subtotal_fill
        cell.font = subtotal_font
        cell.number_format = '#,##0.00'

    # -------------------------------------------------------------
    # CORES DA CONFERÊNCIA
    # -------------------------------------------------------------

    observation_column_index = (
        list(df.columns).index(
            metadata[
                "observation_column"
            ]
        )
        + 1
    )

    for df_index, style_info in row_styles.items():

        # Data começa na linha 3.
        excel_row = (
            df.index.get_loc(
                df_index
            )
            + 3
        )

        if style_info[
            "row_fill"
        ] == "red":

            for column_index in range(
                1,
                ws.max_column + 1,
            ):

                ws.cell(
                    row=excel_row,
                    column=column_index,
                ).fill = red_fill

        else:

            observation_fill = style_info[
                "observation_fill"
            ]

            if observation_fill == "yellow":

                ws.cell(
                    row=excel_row,
                    column=observation_column_index,
                ).fill = yellow_fill

            elif observation_fill == "blue":

                ws.cell(
                    row=excel_row,
                    column=observation_column_index,
                ).fill = blue_fill

    # -------------------------------------------------------------
    # APRESENTAÇÃO
    # -------------------------------------------------------------

    ws.freeze_panes = "A3"

    ws.auto_filter.ref = (
        f"A2:"
        f"{get_column_letter(ws.max_column)}"
        f"{ws.max_row}"
    )

    for column_index in range(
        1,
        ws.max_column + 1,
    ):

        max_length = 0

        for row_index in range(
            1,
            min(
                ws.max_row,
                300,
            )
            + 1,
        ):

            value = ws.cell(
                row=row_index,
                column=column_index,
            ).value

            if value is not None:

                max_length = max(
                    max_length,
                    len(str(value)),
                )

        ws.column_dimensions[
            get_column_letter(
                column_index
            )
        ].width = min(
            max(
                max_length + 2,
                12,
            ),
            55,
        )


def process_accounting_workbook(
    uploaded_accounting_file,
    consolidated_df: pd.DataFrame,
):
    """
    Processa todas as abas visíveis da planilha contábil
    que possuírem um cabeçalho reconhecido.

    Abas ocultas ou sem estrutura contábil são ignoradas.
    """

    consolidated_lookup = (
        prepare_consolidated_lookup(
            consolidated_df
        )
    )

    (
        excel_file,
        hidden_sheets,
    ) = values_only_excel_file(
        uploaded_accounting_file
    )

    processed_sheets = []
    ignored_sheets = []
    warnings = []

    totals = {
        "os_analisadas": 0,
        "conferidas": 0,
        "nao_conferidas": 0,
        "linhas_nao_os": 0,
    }

    for sheet_name in excel_file.sheet_names:

        if sheet_name in hidden_sheets:

            ignored_sheets.append(
                f"{sheet_name} (oculta)"
            )

            continue

        header_row, header_score = (
            detect_accounting_header_row(
                excel_file=excel_file,
                sheet_name=sheet_name,
                max_rows=10,
            )
        )

        if header_row is None:

            ignored_sheets.append(
                f"{sheet_name} (cabeçalho não identificado)"
            )

            continue

        df = excel_file.parse(
            sheet_name=sheet_name,
            header=header_row,
            dtype=object,
            keep_default_na=False,
        )

        # Remove somente colunas completamente sem cabeçalho.
        valid_columns = [
            column
            for column in df.columns
            if normalize_text(column)
            and not normalize_text(
                column
            ).startswith(
                "unnamed"
            )
        ]

        df = df.loc[
            :,
            valid_columns,
        ].copy()

        column_map, ambiguous = (
            build_accounting_column_map(
                df.columns
            )
        )

        required = {
            "HISTORICO",
            "DÉBITO",
            "CRÉDITO",
            "SALDO",
        }

        missing_required = (
            required
            - set(
                column_map.keys()
            )
        )

        ambiguous_required = (
            required
            .intersection(
                ambiguous.keys()
            )
        )

        if (
            missing_required
            or ambiguous_required
        ):

            details = []

            if missing_required:

                details.append(
                    "faltando: "
                    + ", ".join(
                        sorted(
                            missing_required
                        )
                    )
                )

            if ambiguous_required:

                details.append(
                    "duplicadas/ambíguas: "
                    + ", ".join(
                        sorted(
                            ambiguous_required
                        )
                    )
                )

            ignored_sheets.append(
                f"{sheet_name} ({'; '.join(details)})"
            )

            continue

        (
            processed_df,
            row_styles,
            stats,
            metadata,
        ) = process_accounting_sheet(
            df=df,
            column_map=column_map,
            consolidated_lookup=consolidated_lookup,
        )

        processed_sheets.append(
            {
                "sheet_name": sheet_name,
                "df": processed_df,
                "row_styles": row_styles,
                "metadata": metadata,
            }
        )

        for key in totals:
            totals[key] += stats[key]

    if not processed_sheets:

        raise ValueError(
            "Nenhuma aba contábil visível pôde ser processada. "
            "Verifique se HISTORICO, DÉBITO, CRÉDITO e SALDO "
            "estão no cabeçalho das primeiras 10 linhas."
        )

    # -------------------------------------------------------------
    # GERAR EXCEL FINAL
    # -------------------------------------------------------------

    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:

        for item in processed_sheets:

            safe_sheet_name = (
                item[
                    "sheet_name"
                ][:31]
            )

            item[
                "df"
            ].to_excel(
                writer,
                index=False,
                sheet_name=safe_sheet_name,
                startrow=1,
            )

    output.seek(0)

    workbook = load_workbook(
        output
    )

    for item in processed_sheets:

        safe_sheet_name = (
            item[
                "sheet_name"
            ][:31]
        )

        ws = workbook[
            safe_sheet_name
        ]

        style_accounting_output_sheet(
            ws=ws,
            df=item[
                "df"
            ],
            row_styles=item[
                "row_styles"
            ],
            metadata=item[
                "metadata"
            ],
        )

    final_output = io.BytesIO()

    workbook.save(
        final_output
    )

    final_output.seek(0)

    return (
        final_output.getvalue(),
        totals,
        [
            item[
                "sheet_name"
            ]
            for item in processed_sheets
        ],
        ignored_sheets,
        warnings,
    )



# =================================================================
# INTERFACE STREAMLIT
# =================================================================

st.title(
    "📊 Consolidação de Fretes e Conferência Contábil"
)

st.write(
    "O aplicativo possui duas etapas independentes: "
    "consolidação das planilhas de frete e conferência "
    "dos lançamentos enviados pela contabilidade."
)


# =================================================================
# ETAPA 1
# =================================================================

st.header(
    "Etapa 1 — Consolidar fretes"
)

st.caption(
    "Você pode usar somente esta etapa e baixar o consolidado."
)

with st.expander(
    "Regras da consolidação"
):

    st.markdown(
        """
- O cabeçalho é procurado automaticamente nas **10 primeiras linhas** de cada aba.
- Abas **ocultas** não são processadas.
- Abas com **Resumo** no nome são ignoradas.
- Maiúsculas/minúsculas, acentos, espaços e pontuação do cabeçalho são ignorados na comparação.
- `TRANSPORTADOR/MOTORISTA` é tratado como `TRANSPORTADOR`.
- `COMBUSTÍVEL` é tratado como `ABASTECIMENTO`.
- `TOTAL COM DIFERENÇA / DESCONTO` é tratado como `TOTAL COM DESCONTO`.
- `OBSERVAÇÃO` é tratada como `MOTIVO`.
- `FRANQUINA` também é reconhecida como `FRANQUIA`.
- Se uma coluna não existir, ela fica em branco.
- As fórmulas da origem são lidas apenas pelo **resultado salvo**.
- Sem `KM TOTAL`: `FRETE SEM DESCONTO = FRETE SAÍDA + PEDÁGIO + OUTROS CRÉDITOS`.
- Com `KM TOTAL`: `FRETE SEM DESCONTO = FRETE SAÍDA + (ADC (KM × FRANQUIA) + PEDÁGIO + OUTROS CRÉDITOS`.
- `TOTAL COM DESCONTO = FRETE SEM DESCONTO + DESCONTO + ABASTECIMENTO`.
- Se `ABASTECIMENTO` tiver número e `MOTIVO` estiver vazio ou não mencionar abastecimento/combustível, a célula `MOTIVO` é destacada em amarelo.
        """
    )


uploaded_file = st.file_uploader(
    "Selecione a planilha de fretes",
    type=["xlsx"],
    key="freight_upload",
)


if uploaded_file is not None:

    if st.button(
        "Processar consolidação",
        type="primary",
        use_container_width=True,
        key="process_freight",
    ):

        try:

            with st.spinner(
                "Consolidando as abas..."
            ):

                (
                    final_df,
                    report,
                    ignored_sheets,
                    warnings,
                ) = process_workbook(
                    uploaded_file
                )

                excel_bytes = (
                    dataframe_to_excel_bytes(
                        final_df
                    )
                )

            st.session_state[
                "final_df"
            ] = final_df

            st.session_state[
                "report"
            ] = report

            st.session_state[
                "ignored_sheets"
            ] = ignored_sheets

            st.session_state[
                "warnings"
            ] = warnings

            st.session_state[
                "excel_bytes"
            ] = excel_bytes

        except Exception as exc:

            st.error(
                "Não foi possível processar a consolidação. "
                f"Detalhes: {exc}"
            )


if "final_df" in st.session_state:

    final_df = st.session_state[
        "final_df"
    ]

    report = st.session_state.get(
        "report",
        [],
    )

    ignored_sheets = st.session_state.get(
        "ignored_sheets",
        [],
    )

    warnings = st.session_state.get(
        "warnings",
        [],
    )

    excel_bytes = st.session_state[
        "excel_bytes"
    ]

    processed_count = sum(
        1
        for item in report
        if item.get(
            "Status"
        ) == "Processada"
    )

    col1, col2, col3 = st.columns(
        3
    )

    col1.metric(
        "Abas processadas",
        processed_count,
    )

    col2.metric(
        "Registros consolidados",
        len(final_df),
    )

    col3.metric(
        "Abas ignoradas",
        len(ignored_sheets),
    )

    if ignored_sheets:

        st.info(
            "Abas ignoradas: "
            + ", ".join(
                ignored_sheets
            )
        )

    for warning in warnings:
        st.warning(
            warning
        )

    st.download_button(
        label="⬇️ Baixar Resumo Geral.xlsx",
        data=excel_bytes,
        file_name="Resumo_Geral.xlsx",
        mime=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        use_container_width=True,
        key="download_summary",
    )


# =================================================================
# ETAPA 2
# =================================================================

st.divider()

st.header(
    "Etapa 2 — Conferência contábil"
)

st.caption(
    "Você pode usar o consolidado gerado acima ou enviar "
    "um Resumo Geral já criado anteriormente."
)

with st.expander(
    "Regras da conferência contábil"
):

    st.markdown(
        """
- Só são analisadas linhas cujo `HISTORICO` começa com `OS.3` (ou `O.S.3`) seguido da numeração da O.S.
- Linhas que não são O.S. permanecem no arquivo, mas não recebem conferência automática.
- A numeração é procurada na coluna `O.S VIAG` do consolidado.
- Se a O.S. não existir ou aparecer mais de uma vez no consolidado, ela fica `NÃO CONFERIDO`.
- A soma de `DÉBITO` da O.S. deve ser igual a `FRETE SEM DESCONTO`.
- A soma de `CRÉDITO` deve ser igual a `|ABASTECIMENTO| + |DESCONTO|`.
- A soma de `SALDO` deve ser igual a `TOTAL COM DESCONTO`.
- Crédito igual ao valor de `ABASTECIMENTO`: observação de combustível em **amarelo**.
- Crédito igual ao valor de `DESCONTO`: observação de avarias em **azul-claro**.
- Qualquer divergência deixa todas as linhas daquela O.S. em **vermelho-claro** e informa o valor divergente.
- `CONFERÊNCIA` é adicionada como primeira coluna.
- `VENCIMENTO` é preenchido a partir do consolidado.
- A linha acima do cabeçalho recebe os subtotais gerais de `DÉBITO`, `CRÉDITO` e `SALDO`.
        """
    )


if "final_df" in st.session_state:

    consolidated_source_options = [
        "Usar consolidado gerado na Etapa 1",
        "Enviar um consolidado pronto",
    ]

else:

    consolidated_source_options = [
        "Enviar um consolidado pronto",
    ]


consolidated_source = st.radio(
    "Fonte do consolidado",
    consolidated_source_options,
    key="consolidated_source",
)


uploaded_consolidated = None

if (
    consolidated_source
    == "Enviar um consolidado pronto"
):

    uploaded_consolidated = st.file_uploader(
        "Selecione o Resumo Geral consolidado",
        type=["xlsx"],
        key="consolidated_upload",
    )


uploaded_accounting = st.file_uploader(
    "Selecione a planilha da contabilidade",
    type=["xlsx"],
    key="accounting_upload",
)


can_process_accounting = (
    uploaded_accounting is not None
    and (
        consolidated_source
        == "Usar consolidado gerado na Etapa 1"
        or uploaded_consolidated is not None
    )
)


if st.button(
    "Conferir contabilidade",
    type="primary",
    use_container_width=True,
    disabled=not can_process_accounting,
    key="process_accounting",
):

    try:

        with st.spinner(
            "Conferindo as O.S. e os lançamentos contábeis..."
        ):

            if (
                consolidated_source
                == "Usar consolidado gerado na Etapa 1"
            ):

                consolidated_df = (
                    st.session_state[
                        "final_df"
                    ].copy()
                )

            else:

                consolidated_df = (
                    load_consolidated_from_upload(
                        uploaded_consolidated
                    )
                )

            (
                accounting_excel,
                accounting_totals,
                accounting_processed_sheets,
                accounting_ignored_sheets,
                accounting_warnings,
            ) = process_accounting_workbook(
                uploaded_accounting_file=uploaded_accounting,
                consolidated_df=consolidated_df,
            )

        st.session_state[
            "accounting_excel"
        ] = accounting_excel

        st.session_state[
            "accounting_totals"
        ] = accounting_totals

        st.session_state[
            "accounting_processed_sheets"
        ] = accounting_processed_sheets

        st.session_state[
            "accounting_ignored_sheets"
        ] = accounting_ignored_sheets

        st.session_state[
            "accounting_warnings"
        ] = accounting_warnings

    except Exception as exc:

        st.error(
            "Não foi possível realizar a conferência contábil. "
            f"Detalhes: {exc}"
        )


if "accounting_excel" in st.session_state:

    accounting_totals = (
        st.session_state[
            "accounting_totals"
        ]
    )

    processed_sheets = (
        st.session_state[
            "accounting_processed_sheets"
        ]
    )

    ignored_sheets = (
        st.session_state[
            "accounting_ignored_sheets"
        ]
    )

    warnings = (
        st.session_state[
            "accounting_warnings"
        ]
    )

    col1, col2, col3, col4 = st.columns(
        4
    )

    col1.metric(
        "O.S. analisadas",
        accounting_totals[
            "os_analisadas"
        ],
    )

    col2.metric(
        "Conferidas",
        accounting_totals[
            "conferidas"
        ],
    )

    col3.metric(
        "Não conferidas",
        accounting_totals[
            "nao_conferidas"
        ],
    )

    col4.metric(
        "Linhas não-O.S.",
        accounting_totals[
            "linhas_nao_os"
        ],
    )

    st.success(
        "Abas contábeis processadas: "
        + ", ".join(
            processed_sheets
        )
    )

    if ignored_sheets:

        st.info(
            "Abas contábeis ignoradas: "
            + ", ".join(
                ignored_sheets
            )
        )

    for warning in warnings:
        st.warning(
            warning
        )

    st.download_button(
        label="⬇️ Baixar Contabilidade Conferida.xlsx",
        data=st.session_state[
            "accounting_excel"
        ],
        file_name="Contabilidade_Conferida.xlsx",
        mime=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        use_container_width=True,
        key="download_accounting",
    )
